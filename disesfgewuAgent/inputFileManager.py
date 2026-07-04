import asyncio
import os
import struct
import tarfile
import tempfile
import zipfile
from pathlib import Path
from typing import Callable, Optional, Tuple
from pypdf import PdfReader
from pptx import Presentation
from openpyxl import load_workbook
from docx import Document


class inputFileManager:

    SUPPORTED_FORMATS = {
        ".txt", ".md", ".pdf", ".xlsx", ".docx", ".pptx",
        ".jpg", ".jpeg", ".png", ".zip", ".tar", ".gz", ".tgz",
    }
    ARCHIVE_EXTRACT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".py", ".js", ".ts", ".html", ".css", ".xml", ".pdf", ".xlsx", ".docx", ".pptx"}
    MAX_ARCHIVE_MEMBERS = 50
    MAX_ARCHIVE_MEMBER_BYTES = 10 * 1024 * 1024
    MAX_ARCHIVE_TOTAL_BYTES = 50 * 1024 * 1024

    @staticmethod
    def extract(filepath: str) -> str:
        path = Path(filepath)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {filepath}")

        ext = path.suffix.lower()
        suffixes = [s.lower() for s in path.suffixes]
        if ext == ".pdf":
            return inputFileManager._extractPdf(filepath)
        elif ext == ".xlsx":
            return inputFileManager._extractXlsx(filepath)
        elif ext == ".docx":
            return inputFileManager._extractDocx(filepath)
        elif ext == ".pptx":
            return inputFileManager._extractPptx(filepath)
        elif ext in {".jpg", ".jpeg", ".png"}:
            return inputFileManager._extractImageMetadata(filepath)
        elif ext == ".zip":
            return inputFileManager._extractZip(filepath)
        elif ext in {".tar", ".tgz"} or suffixes[-2:] == [".tar", ".gz"]:
            return inputFileManager._extractTar(filepath)
        # Everything else (.txt, .md, source code, .json, .csv, .yaml, ...) is
        # read as plain text. _extractTxt falls back to latin-1, so it won't
        # crash even on non-UTF-8 text-like files fed in by mistake.
        return inputFileManager._extractTxt(filepath)

    @staticmethod
    def _extractTxt(filepath: str) -> str:
        try:
            with open(filepath, "r", encoding="utf-8") as f:
                return f.read()
        except UnicodeDecodeError:
            # latin-1 maps every byte, so this never raises and preserves
            # the original bytes rather than crashing on non-UTF-8 input.
            with open(filepath, "r", encoding="latin-1") as f:
                return f.read()

    @staticmethod
    def _extractPdf(filepath: str) -> str:

        reader = PdfReader(filepath)
        pages = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                pages.append(f"[Page {i + 1}]\n{text}")
        return "\n\n".join(pages) if pages else ""

    @staticmethod
    def _extractXlsx(filepath: str) -> str:

        wb = load_workbook(filepath, read_only=True, data_only=True)
        sheets = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(cell) if cell is not None else "" for cell in row]
                if any(cells):
                    rows.append(", ".join(cells))
            if rows:
                sheets.append(f"[Sheet: {sheet_name}]\n" + "\n".join(rows))

        wb.close()
        return "\n\n".join(sheets) if sheets else ""

    @staticmethod
    def _extractDocx(filepath: str) -> str:

        doc = Document(filepath)
        parts = []

        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text)

        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    table_rows.append(" | ".join(cells))
            if table_rows:
                parts.append("[Table]\n" + "\n".join(table_rows))

        return "\n\n".join(parts) if parts else ""

    @staticmethod
    def _extractPptx(filepath: str) -> str:
        prs = Presentation(filepath)
        slides = []

        for i, slide in enumerate(prs.slides):
            slide_parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_parts.append(shape.text.strip())
            if slide_parts:
                slides.append(f"[Slide {i + 1}]\n" + "\n".join(slide_parts))

        return "\n\n".join(slides) if slides else ""

    @staticmethod
    def _extractImageMetadata(filepath: str) -> str:
        path = Path(filepath)
        ext = path.suffix.lower()
        size = path.stat().st_size
        dimensions = inputFileManager._readImageDimensions(filepath)
        dim_text = f"{dimensions[0]}x{dimensions[1]}" if dimensions else "unknown"
        return (
            f"[Image]\n"
            f"format: {ext.lstrip('.').upper()}\n"
            f"filename: {path.name}\n"
            f"bytes: {size}\n"
            f"dimensions: {dim_text}\n"
            "text_extraction: not available (no OCR is performed by inputFileManager)"
        )

    @staticmethod
    def _readImageDimensions(filepath: str) -> Optional[Tuple[int, int]]:
        with open(filepath, "rb") as f:
            header = f.read(32)
            if header.startswith(b"\x89PNG\r\n\x1a\n") and len(header) >= 24:
                return struct.unpack(">II", header[16:24])
            if header.startswith(b"\xff\xd8"):
                f.seek(2)
                while True:
                    marker_start = f.read(1)
                    if not marker_start:
                        return None
                    if marker_start != b"\xff":
                        continue
                    marker = f.read(1)
                    while marker == b"\xff":
                        marker = f.read(1)
                    if not marker or marker in {b"\xd8", b"\xd9"}:
                        continue
                    length_bytes = f.read(2)
                    if len(length_bytes) != 2:
                        return None
                    length = struct.unpack(">H", length_bytes)[0]
                    if marker in {b"\xc0", b"\xc1", b"\xc2", b"\xc3", b"\xc5", b"\xc6", b"\xc7", b"\xc9", b"\xca", b"\xcb", b"\xcd", b"\xce", b"\xcf"}:
                        data = f.read(5)
                        if len(data) != 5:
                            return None
                        height, width = struct.unpack(">HH", data[1:5])
                        return width, height
                    f.seek(max(length - 2, 0), os.SEEK_CUR)
        return None

    @staticmethod
    def _extractZip(filepath: str) -> str:
        with zipfile.ZipFile(filepath) as zf:
            entries = []
            total = 0
            for info in zf.infolist():
                if info.is_dir():
                    continue
                if len(entries) >= inputFileManager.MAX_ARCHIVE_MEMBERS:
                    entries.append("[Archive notice]\nmember limit reached; remaining files skipped")
                    break
                if info.file_size > inputFileManager.MAX_ARCHIVE_MEMBER_BYTES:
                    entries.append(f"[Archive member: {info.filename}]\nskipped: member too large ({info.file_size} bytes)")
                    continue
                total += info.file_size
                if total > inputFileManager.MAX_ARCHIVE_TOTAL_BYTES:
                    entries.append("[Archive notice]\ntotal extraction byte limit reached; remaining files skipped")
                    break
                data = zf.read(info)
                entries.append(inputFileManager._extractArchiveMember(info.filename, data))
            return "[Archive: zip]\n" + "\n\n".join(entries)

    @staticmethod
    def _extractTar(filepath: str) -> str:
        with tarfile.open(filepath) as tf:
            entries = []
            total = 0
            for member in tf.getmembers():
                if not member.isfile():
                    continue
                if len(entries) >= inputFileManager.MAX_ARCHIVE_MEMBERS:
                    entries.append("[Archive notice]\nmember limit reached; remaining files skipped")
                    break
                if member.size > inputFileManager.MAX_ARCHIVE_MEMBER_BYTES:
                    entries.append(f"[Archive member: {member.name}]\nskipped: member too large ({member.size} bytes)")
                    continue
                total += member.size
                if total > inputFileManager.MAX_ARCHIVE_TOTAL_BYTES:
                    entries.append("[Archive notice]\ntotal extraction byte limit reached; remaining files skipped")
                    break
                stream = tf.extractfile(member)
                if stream is None:
                    continue
                entries.append(inputFileManager._extractArchiveMember(member.name, stream.read()))
            suffixes = "".join(Path(filepath).suffixes[-2:]).lower()
            kind = "tar.gz" if suffixes == ".tar.gz" or Path(filepath).suffix.lower() == ".tgz" else "tar"
            return f"[Archive: {kind}]\n" + "\n\n".join(entries)

    @staticmethod
    def _extractArchiveMember(name: str, data: bytes) -> str:
        safe_name = name.replace("\\", "/")
        ext = Path(safe_name).suffix.lower()
        if ext not in inputFileManager.ARCHIVE_EXTRACT_EXTENSIONS:
            return f"[Archive member: {safe_name}]\nskipped: unsupported member type"
        try:
            text = inputFileManager._extractBytesWithSuffix(data, ext)
        except Exception as exc:
            return f"[Archive member: {safe_name}]\nskipped: extraction failed ({exc})"
        return f"[Archive member: {safe_name}]\n{text}"

    @staticmethod
    def _extractBytesWithSuffix(data: bytes, suffix: str) -> str:
        if suffix in {".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".py", ".js", ".ts", ".html", ".css", ".xml"}:
            try:
                return data.decode("utf-8")
            except UnicodeDecodeError:
                return data.decode("latin-1")
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            tmp.write(data)
            tmp_path = tmp.name
        try:
            if suffix == ".pdf":
                return inputFileManager._extractPdf(tmp_path)
            if suffix == ".xlsx":
                return inputFileManager._extractXlsx(tmp_path)
            if suffix == ".docx":
                return inputFileManager._extractDocx(tmp_path)
            if suffix == ".pptx":
                return inputFileManager._extractPptx(tmp_path)
            return ""
        finally:
            try:
                os.remove(tmp_path)
            except OSError:
                pass

    @staticmethod
    async def decompose(filepaths: list, countTokens: Callable[[str], int]) -> tuple:
        parts = []
        for filepath in filepaths:
            text = await asyncio.to_thread(inputFileManager.extract, filepath)
            filename = os.path.basename(filepath)
            abspath = os.path.abspath(filepath)
            parts.append(f"[{filename}] (path: {abspath})\n{text}")

        result = "[FILES]\n" + "\n\n".join(parts)
        token_count = countTokens(result)
        return result, token_count

    @staticmethod
    def decomposeSync(filepaths: list, countTokens: Callable[[str], int]) -> tuple:
        parts = []
        for filepath in filepaths:
            text = inputFileManager.extract(filepath)
            filename = os.path.basename(filepath)
            abspath = os.path.abspath(filepath)
            parts.append(f"[{filename}] (path: {abspath})\n{text}")

        result = "[FILES]\n" + "\n\n".join(parts)
        token_count = countTokens(result)
        return result, token_count