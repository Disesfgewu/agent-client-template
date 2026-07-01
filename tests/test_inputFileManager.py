import unittest
import asyncio
import tempfile
import shutil
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from disesfgewuAgent.inputFileManager import inputFileManager


class TestInputFileManagerTxt(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.test_dir)

    def test_extract_txt(self):
        filepath = os.path.join(self.test_dir, "test.txt")
        content = "Hello, this is a test file.\nLine 2."
        with open(filepath, "w", encoding="utf-8") as f:
            f.write(content)

        result = inputFileManager.extract(filepath)
        self.assertEqual(result, content)

    def test_extract_md(self):
        filepath = os.path.join(self.test_dir, "test.md")
        content = "# Heading\n\nSome **markdown** content.\n\n- Item 1\n- Item 2"
        with open(filepath, "w", encoding="utf-8") as f:
            f.write(content)

        result = inputFileManager.extract(filepath)
        self.assertEqual(result, content)

    def test_extract_txt_empty(self):
        filepath = os.path.join(self.test_dir, "empty.txt")
        with open(filepath, "w", encoding="utf-8") as f:
            f.write("")

        result = inputFileManager.extract(filepath)
        self.assertEqual(result, "")

    def test_extract_txt_unicode(self):
        filepath = os.path.join(self.test_dir, "unicode.txt")
        content = "Hello 你好 こんにちは 🎉"
        with open(filepath, "w", encoding="utf-8") as f:
            f.write(content)

        result = inputFileManager.extract(filepath)
        self.assertEqual(result, content)


class TestInputFileManagerPdf(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.test_dir)

    def test_extract_pdf(self):
        from pypdf import PdfWriter

        filepath = os.path.join(self.test_dir, "test.pdf")
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)

        with open(filepath, "wb") as f:
            writer.write(f)

        result = inputFileManager.extract(filepath)
        self.assertIsInstance(result, str)


class TestInputFileManagerXlsx(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.test_dir)

    def test_extract_xlsx(self):
        from openpyxl import Workbook

        filepath = os.path.join(self.test_dir, "test.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "Name"
        ws["B1"] = "Age"
        ws["A2"] = "Alice"
        ws["B2"] = 30
        ws["A3"] = "Bob"
        ws["B3"] = 25
        wb.save(filepath)

        result = inputFileManager.extract(filepath)

        self.assertIn("[Sheet: Sheet1]", result)
        self.assertIn("Name", result)
        self.assertIn("Age", result)
        self.assertIn("Alice", result)
        self.assertIn("30", result)

    def test_extract_xlsx_multiple_sheets(self):
        from openpyxl import Workbook

        filepath = os.path.join(self.test_dir, "multi.xlsx")
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Data"
        ws1["A1"] = "Value1"

        ws2 = wb.create_sheet("Extra")
        ws2["A1"] = "Value2"

        wb.save(filepath)

        result = inputFileManager.extract(filepath)

        self.assertIn("[Sheet: Data]", result)
        self.assertIn("[Sheet: Extra]", result)
        self.assertIn("Value1", result)
        self.assertIn("Value2", result)

    def test_extract_xlsx_empty(self):
        from openpyxl import Workbook

        filepath = os.path.join(self.test_dir, "empty.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Empty"
        wb.save(filepath)

        result = inputFileManager.extract(filepath)
        self.assertEqual(result, "")


class TestInputFileManagerDocx(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.test_dir)

    def test_extract_docx(self):
        from docx import Document

        filepath = os.path.join(self.test_dir, "test.docx")
        doc = Document()
        doc.add_paragraph("First paragraph.")
        doc.add_paragraph("Second paragraph.")
        doc.save(filepath)

        result = inputFileManager.extract(filepath)

        self.assertIn("First paragraph.", result)
        self.assertIn("Second paragraph.", result)

    def test_extract_docx_with_table(self):
        from docx import Document

        filepath = os.path.join(self.test_dir, "table.docx")
        doc = Document()
        doc.add_paragraph("Before table.")

        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Header1"
        table.cell(0, 1).text = "Header2"
        table.cell(1, 0).text = "Data1"
        table.cell(1, 1).text = "Data2"

        doc.save(filepath)

        result = inputFileManager.extract(filepath)

        self.assertIn("Before table.", result)
        self.assertIn("[Table]", result)
        self.assertIn("Header1", result)
        self.assertIn("Data1", result)

    def test_extract_docx_empty(self):
        from docx import Document

        filepath = os.path.join(self.test_dir, "empty.docx")
        doc = Document()
        doc.save(filepath)

        result = inputFileManager.extract(filepath)
        self.assertEqual(result, "")


class TestInputFileManagerPptx(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.test_dir)

    def test_extract_pptx(self):
        from pptx import Presentation

        filepath = os.path.join(self.test_dir, "test.pptx")
        prs = Presentation()

        slide1 = prs.slides.add_slide(prs.slide_layouts[0])
        slide1.shapes.title.text = "Slide 1 Title"
        slide1.placeholders[1].text = "Subtitle 1"

        slide2 = prs.slides.add_slide(prs.slide_layouts[0])
        slide2.shapes.title.text = "Slide 2 Title"

        prs.save(filepath)

        result = inputFileManager.extract(filepath)

        self.assertIn("[Slide 1]", result)
        self.assertIn("Slide 1 Title", result)
        self.assertIn("[Slide 2]", result)
        self.assertIn("Slide 2 Title", result)

    def test_extract_pptx_empty(self):
        from pptx import Presentation

        filepath = os.path.join(self.test_dir, "empty.pptx")
        prs = Presentation()
        prs.save(filepath)

        result = inputFileManager.extract(filepath)
        self.assertEqual(result, "")


class TestInputFileManagerErrors(unittest.TestCase):
    def test_extract_file_not_found(self):
        with self.assertRaises(FileNotFoundError) as context:
            inputFileManager.extract("/nonexistent/path/file.txt")
        self.assertIn("File not found", str(context.exception))

    def test_extract_unknown_extension_reads_as_text(self):
        test_dir = tempfile.mkdtemp()
        try:
            filepath = os.path.join(test_dir, "file.xyz")
            with open(filepath, "w", encoding="utf-8") as f:
                f.write("content")

            self.assertEqual(inputFileManager.extract(filepath), "content")
        finally:
            shutil.rmtree(test_dir)

    def test_extract_source_code_as_text(self):
        test_dir = tempfile.mkdtemp()
        try:
            filepath = os.path.join(test_dir, "script.py")
            with open(filepath, "w", encoding="utf-8") as f:
                f.write("def hello():\n    return 42\n")

            result = inputFileManager.extract(filepath)
            self.assertIn("def hello", result)
            self.assertIn("return 42", result)
        finally:
            shutil.rmtree(test_dir)


class TestInputFileManagerDecompose(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()

        self.file1 = os.path.join(self.test_dir, "file1.txt")
        with open(self.file1, "w", encoding="utf-8") as f:
            f.write("Content of file 1.")

        self.file2 = os.path.join(self.test_dir, "file2.txt")
        with open(self.file2, "w", encoding="utf-8") as f:
            f.write("Content of file 2.")

    def tearDown(self):
        shutil.rmtree(self.test_dir)

    def _countTokens(self, text: str) -> int:
        return len(text.split())

    def test_decomposeSync(self):
        result, tokens = inputFileManager.decomposeSync(
            [self.file1, self.file2], self._countTokens
        )

        self.assertIn("[FILES]", result)
        self.assertIn("[file1.txt]", result)
        self.assertIn("[file2.txt]", result)
        self.assertIn("Content of file 1.", result)
        self.assertIn("Content of file 2.", result)
        self.assertIsInstance(tokens, int)
        self.assertGreater(tokens, 0)

    def test_decomposeSync_empty_list(self):
        result, tokens = inputFileManager.decomposeSync([], self._countTokens)

        self.assertEqual(result, "[FILES]\n")
        self.assertEqual(tokens, 1)


class TestInputFileManagerDecomposeAsync(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()

        self.file1 = os.path.join(self.test_dir, "async1.txt")
        with open(self.file1, "w", encoding="utf-8") as f:
            f.write("Async content 1.")

        self.file2 = os.path.join(self.test_dir, "async2.txt")
        with open(self.file2, "w", encoding="utf-8") as f:
            f.write("Async content 2.")

    def tearDown(self):
        shutil.rmtree(self.test_dir)

    def _countTokens(self, text: str) -> int:
        return len(text.split())

    async def test_decompose(self):
        result, tokens = await inputFileManager.decompose(
            [self.file1, self.file2], self._countTokens
        )

        self.assertIn("[FILES]", result)
        self.assertIn("[async1.txt]", result)
        self.assertIn("[async2.txt]", result)
        self.assertIn("Async content 1.", result)
        self.assertIn("Async content 2.", result)
        self.assertIsInstance(tokens, int)
        self.assertGreater(tokens, 0)

    async def test_decompose_empty_list(self):
        result, tokens = await inputFileManager.decompose([], self._countTokens)

        self.assertEqual(result, "[FILES]\n")
        self.assertEqual(tokens, 1)

    async def test_decompose_mixed_formats(self):
        from openpyxl import Workbook

        xlsx_path = os.path.join(self.test_dir, "data.xlsx")
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "TestData"
        wb.save(xlsx_path)

        result, tokens = await inputFileManager.decompose(
            [self.file1, xlsx_path], self._countTokens
        )

        self.assertIn("[async1.txt]", result)
        self.assertIn("[data.xlsx]", result)
        self.assertIn("Async content 1.", result)
        self.assertIn("TestData", result)


if __name__ == "__main__":
    unittest.main()
